#!/usr/bin/env python3
"""
Excel workbook for checking HeatAQ's forecasts by hand.

For a weather station and two days, the workbook shows every step from MET
Norway's raw forecast values to the forecast HeatAQ stores and plans with:

- Grid: the station's position on the MET Nordic grid (projection formulas),
  the 4 grid points around it, their weights, and the height correction;
- Runs: the forecast runs and when they became available, with links to the
  raw data in MET's archive (a reader can open them in a browser);
- Forecast: each forecast hour recomputed by formulas from the 4 grid points'
  raw values, next to HeatAQ's own value (OK when equal within rounding);
- Night plan: which run the predictive controller plans the night with, and
  when it plans again.

Every computed cell is a formula. Raw values are fetched here, independently
of HeatAQ (fetch_met_nordic.py's code); HeatAQ's values come from running
the app's own code (check_app_values.php, needs php with curl).

    python check_workbook.py --name "Røros lufthavn" --lat 62.5781 --lon 11.3487 \\
        --elevation 625 --date 2024-01-04 [--close 21:00 --open 07:00] [--out file.xlsx]

Needs openpyxl (pip install openpyxl). Excel computes the formulas when the
file is opened; in LibreOffice use Data -> Calculate -> Recalculate Hard.
See README.md in this folder.
"""

import argparse
import datetime as dt
import json
import os
import re
import subprocess
import sys
from zoneinfo import ZoneInfo

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import fetch_met_nordic as mn  # noqa: E402  (independent fetch, parsing and projection)

from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402
from openpyxl.workbook.defined_name import DefinedName  # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, '..', '..'))
LEAD_HOURS = 36          # HeatAQ reads hours 0-36 of each run
LAND_THRESHOLD = 0.5

F = 'Arial'
f_norm = Font(name=F, size=10)
f_bold = Font(name=F, size=10, bold=True)
f_title = Font(name=F, size=14, bold=True)
f_h2 = Font(name=F, size=11, bold=True)
f_in = Font(name=F, size=10, color='0000FF')            # from MET's archive, or an input
f_app = Font(name=F, size=10, color='008000')           # computed by HeatAQ
f_link = Font(name=F, size=10, color='0563C1', underline='single')
f_note = Font(name=F, size=9, italic=True, color='666666')
fill_head = PatternFill('solid', fgColor='DDE4EE')
fill_in = PatternFill('solid', fgColor='FFF2CC')        # inputs a reader may change
box = Border(bottom=Side(style='thin', color='B7B7B7'))
wrap = Alignment(wrap_text=True, vertical='top')
center = Alignment(horizontal='center', vertical='center', wrap_text=True)
DT = 'yyyy-mm-dd hh:mm'
DTS = 'yyyy-mm-dd hh:mm:ss'


def log(msg):
    print(msg, file=sys.stderr, flush=True)


def ts(s):
    return dt.datetime.strptime(s, '%Y-%m-%d %H:%M:%S')


# --------------------------------------------------------------------------
# Data
# --------------------------------------------------------------------------

def run_path(day, hh):
    return f"{day:%Y/%m/%d}/met_forecast_1_0km_nordic_{day:%Y%m%d}T{hh:02d}Z.nc"


def fetch_raw(lat, lon, days):
    """Raw values from MET's archive, independently of HeatAQ: grid,
    projection, the 4 grid points, and hours 0-36 of every run."""
    sample = None
    for day in days:
        for hh in (0, 6, 12, 18):
            if mn.http_get(mn.dap_url(run_path(day, hh), 'dds')) is not None:
                sample = run_path(day, hh)
                break
        if sample:
            break
    if sample is None:
        raise RuntimeError('No forecast file in the archive for these days')
    das = mn.http_get(mn.dap_url(sample, 'das'))
    dds = mn.http_get(mn.dap_url(sample, 'dds'))
    xy = mn.parse_ascii(mn.http_get(mn.dap_url(sample, 'ascii', 'x[0:1],y[0:1]')))
    nx = int(re.search(r'x\[x = (\d+)\]', dds).group(1))
    ny = int(re.search(r'y\[y = (\d+)\]', dds).group(1))
    grid = mn.Grid(das, xy['x'][0], xy['x'][1] - xy['x'][0], xy['y'][0], xy['y'][1] - xy['y'][0], nx, ny)
    fi, fj = grid.fractional_index(lat, lon)
    i, j = int(fi // 1), int(fj // 1)
    if not (0 <= i < nx - 1 and 0 <= j < ny - 1):
        raise RuntimeError('The station is outside the MET Nordic grid (Nordic countries only)')
    sl = f'[{j}:{j + 1}][{i}:{i + 1}]'
    corners_url = mn.dap_url(sample, 'ascii', f'latitude{sl},longitude{sl},altitude{sl},land_area_fraction{sl}')
    block = re.search(r'projection_lcc \{(.*?)\}', das, re.S).group(1)
    proj = {k: float(re.search(r'Float64 ' + k + r' ([-\d.eE+]+)', block).group(1))
            for k in ('standard_parallel', 'latitude_of_projection_origin', 'longitude_of_central_meridian', 'earth_radius')}
    raw = {
        'i': i, 'j': j, 'sample': sample, 'proj': proj, 'x': xy['x'], 'y': xy['y'],
        'corners': mn.parse_ascii(mn.http_get(corners_url)),
        'links': {'sample_das': mn.dap_url(sample, 'das'),
                  'sample_xy': mn.dap_url(sample, 'ascii', 'x[0:1],y[0:1]'),
                  'sample_corners': corners_url},
        'runs': [], 'missing': [],
    }
    for day in days:
        for hh in (0, 6, 12, 18):
            path = run_path(day, hh)
            issued = f'{day:%Y-%m-%d} {hh:02d}:00:00'
            run_dds = mn.http_get(mn.dap_url(path, 'dds'))
            run_das = mn.http_get(mn.dap_url(path, 'das'))
            if run_dds is None or run_das is None:
                raw['missing'].append(issued)
                log(f'  {issued} not in the archive')
                continue
            steps = int(re.search(r'time\[time = (\d+)\]', run_dds).group(1))
            last = min(LEAD_HOURS, steps - 1)
            sl3 = f'[0:{last}]{sl}'
            data_url = mn.dap_url(path, 'ascii', f'time[0:{last}],air_temperature_2m{sl3},'
                                                 f'wind_speed_10m{sl3},relative_humidity_2m{sl3}')
            data = mn.parse_ascii(mn.http_get(data_url))
            m = re.search(r'String history "Created (\d{4}-\d{2}-\d{2})T(\d{2}:\d{2}:\d{2})Z"', run_das)
            raw['runs'].append({
                'issued': issued, 'created': f'{m.group(1)} {m.group(2)}' if m else None,
                'data': data_url, 'das': mn.dap_url(path, 'das'),
                'html': f'{mn.THREDDS}/dodsC/{mn.ARCHIVE}/{path}.html',
                'catalog': f'{mn.THREDDS}/catalog/{mn.ARCHIVE}/{day:%Y/%m/%d}/catalog.html',
                'hours': len(data['time']),
                'air': data['air_temperature_2m'], 'wind': data['wind_speed_10m'], 'rh': data['relative_humidity_2m'],
            })
            log(f'  {issued}: {len(data["time"])} hours')
    return raw


def app_values(php, lat, lon, elevation, date):
    """HeatAQ's own values: lib/MetNordicForecast.php run by check_app_values.php."""
    cmd = [php, os.path.join(HERE, 'check_app_values.php'), f'--lat={lat}', f'--lon={lon}',
           f'--elevation={elevation}', f'--date={date}', '--days=2']
    res = subprocess.run(cmd, capture_output=True, text=True)
    if res.returncode != 0:
        raise RuntimeError(f'check_app_values.php failed: {res.stderr.strip() or res.stdout.strip()}')
    return json.loads(res.stdout)


# --------------------------------------------------------------------------
# Workbook
# --------------------------------------------------------------------------

def put(ws, ref, value, font=f_norm, fmt=None, fill=None, align=None):
    c = ws[ref]
    c.value = value
    c.font = font
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    if align:
        c.alignment = align
    return c


def link(ws, ref, text, url, fmt=None):
    c = put(ws, ref, text, f_link, fmt)
    c.hyperlink = url
    return c


def header(ws, row, labels, col=1):
    for k, lab in enumerate(labels):
        c = ws.cell(row=row, column=col + k, value=lab)
        c.font = f_bold
        c.fill = fill_head
        c.alignment = center


def build(args, raw, app, utc_offset, out):
    wb = Workbook()
    wb.calculation.fullCalcOnLoad = True
    d0 = dt.date.fromisoformat(args.date)
    d1 = d0 + dt.timedelta(days=1)
    try:
        version = json.load(open(os.path.join(ROOT, 'assets', 'js', 'sim', 'version.json'))).get('version', '')
    except OSError:
        version = ''

    # ------------------------------------------------------------ Read me
    rm = wb.active
    rm.title = 'Read me'
    rm.column_dimensions['A'].width = 120
    lines = [
        (f'HeatAQ forecast check: {args.name}, {d0:%d} and {d1:%d %B %Y}', f_title),
        ('', None),
        ('What this workbook checks', f_h2),
        ('1. Grid: which 4 MET Nordic grid points surround the weather station, and their weights (sheet Grid).', None),
        ('2. Forecast: each forecast value HeatAQ stores, recomputed here from the raw values in MET\'s archive '
         '(sheet Forecast). Column "Check" says OK when HeatAQ\'s value equals this sheet\'s within rounding.', None),
        ('3. Runs: when each forecast run became available (sheet Runs).', None),
        ('4. Night plan: which forecast run the predictive controller plans the night with, and when it plans '
         'again (sheet Night plan).', None),
        ('', None),
        ('How to check', f_h2),
        ('- Every raw number (blue) comes from MET Norway\'s archive. Click the links (sheet Runs, and the first '
         'hour of each run in sheet Forecast) to open the exact request in a browser and compare the numbers. '
         'The 4 values of each hour are listed south-west, south-east, north-west, north-east.', None),
        ('- Every computed number (black) is a formula: click a cell to see how it is computed.', None),
        ('- Green numbers are HeatAQ\'s own results, for comparison. The differences should be rounding only '
         '(HeatAQ stores air and wind with 2 decimals, humidity with 1).', None),
        ('- Yellow cells are inputs you can change (e.g. the station height or the pool\'s opening hours); the '
         'formulas follow. HeatAQ\'s values (green) were computed with the values shown, so they then differ.', None),
        ('', None),
        ('Colours', f_h2),
        ('Blue: from MET\'s archive (or an input).   Black: formula.   Green: computed by HeatAQ.   '
         'Yellow fill: input you may change.', None),
        ('', None),
        ('Method (as in HeatAQ: lib/MetNordicForecast.php, sim-ts/src/engine.ts)', f_h2),
        ('- MET Nordic forecasts: 1 km grid over the Nordic countries, runs at 00, 06, 12 and 18 UTC, hourly '
         'values; HeatAQ reads hours 0-36 of each run.', None),
        ('- The station\'s position is projected onto the grid (Lambert conformal conic, parameters from the '
         'file). The 4 surrounding grid points are weighted bilinearly (all on land, or all at sea), otherwise '
         'by inverse distance over the land points only.', None),
        ('- Air temperature is moved from the grid\'s height to the station\'s height at 0.65 °C per 100 m. Wind '
         'is not changed (negative values set to 0); humidity is converted to % and kept within 0-100.', None),
        ('- A run is usable from its file\'s creation time ("Created" in the file metadata), or 1 hour after the '
         'run time if that is missing or implausible.', None),
        ('- Planning: when the pool closes, the controller plans the night with the newest run available at that '
         'hour (and covering it). During the closed period it plans again whenever a newer run becomes '
         'available. Hours beyond a run\'s last hour use its last hour.', None),
        ('', None),
        ('Data', f_h2),
        ('MET Nordic forecast archive, MET Norway (thredds.met.no, metpparchive). Licence: CC BY 4.0, credit '
         'MET Norway.', None),
        (f'Station: {args.name}, {args.lat}, {args.lon}, height {args.elevation:g} m.', None),
        ((f'Runs not in MET\'s archive: {", ".join(raw["missing"])}.' if raw['missing'] else
          'All 8 runs of the two days are in MET\'s archive.'), None),
        (f'Made {dt.date.today().isoformat()} with scripts/forecast/check_workbook.py'
         + (f', HeatAQ simulator {version}' if version else '') + '.', f_note),
    ]
    for k, (text, font) in enumerate(lines, start=1):
        c = rm.cell(row=k, column=1, value=text)
        c.font = font or f_norm
        c.alignment = Alignment(wrap_text=True, vertical='top')
    link(rm, f'A{len(lines) + 2}', 'MET Nordic archive (catalog)',
         f'{mn.THREDDS}/catalog/{mn.ARCHIVE}/catalog.html')

    # ------------------------------------------------------------ Grid
    g = wb.create_sheet('Grid')
    for col, w in zip('ABCDEFGHIJKLMNOP', [46, 14, 14, 6, 6, 11, 11, 10, 10, 8, 11, 11, 11, 11, 11, 11]):
        g.column_dimensions[col].width = w
    put(g, 'A1', 'Station position on the MET Nordic grid, and the 4 grid points', f_title)
    put(g, 'A3', 'Station and settings (inputs)', f_h2)
    for r, lab, val, fmt, note in [
        (4, 'Latitude (°)', args.lat, '0.0000', f'{args.name}, as in HeatAQ (Admin → Weather)'),
        (5, 'Longitude (°)', args.lon, '0.0000', ''),
        (6, 'Station height (m)', args.elevation, '0', 'As in HeatAQ (Admin → Weather)'),
        (7, 'Lapse rate (°C per m)', app['lapse'], '0.0000', 'Standard atmosphere (MetNordicForecast::LAPSE)'),
        (8, 'Local time = UTC + (h)', utc_offset, '0', f'{args.tz} on {d0:%d %B %Y}'),
    ]:
        put(g, f'A{r}', lab)
        put(g, f'B{r}', val, f_in, fmt, fill_in)
        put(g, f'C{r}', note, f_note)

    put(g, 'A10', 'MET Nordic grid (from the file metadata)', f_h2)
    link(g, 'C10', 'Open file metadata ↗', raw['links']['sample_das'])
    p = raw['proj']
    for r, lab, val, fmt in [
        (11, 'Standard parallel (°)', p['standard_parallel'], '0.0'),
        (12, 'Latitude of origin (°)', p['latitude_of_projection_origin'], '0.0'),
        (13, 'Central meridian (°)', p['longitude_of_central_meridian'], '0.0'),
        (14, 'Earth radius (m)', p['earth_radius'], '0'),
        (15, 'x of grid column 0 (m)', raw['x'][0], '0.0'),
        (16, 'x of grid column 1 (m)', raw['x'][1], '0.0'),
        (17, 'y of grid row 0 (m)', raw['y'][0], '0.0'),
        (18, 'y of grid row 1 (m)', raw['y'][1], '0.0'),
    ]:
        put(g, f'A{r}', lab)
        put(g, f'B{r}', val, f_in, fmt)
    link(g, 'C15', 'Open x, y ↗', raw['links']['sample_xy'])

    put(g, 'A20', 'Station on the grid (Lambert conformal conic)', f_h2)
    for r, lab, f, fmt in [
        (21, 'n', '=SIN(RADIANS(B11))', '0.000000'),
        (22, 'F', '=COS(RADIANS(B11))*TAN(PI()/4+RADIANS(B11)/2)^B21/B21', '0.000000'),
        (23, 'ρ0 (m)', '=B14*B22/TAN(PI()/4+RADIANS(B12)/2)^B21', '0.0'),
        (24, 'ρ at the station (m)', '=B14*B22/TAN(PI()/4+RADIANS(B4)/2)^B21', '0.0'),
        (25, 'θ (radians)', '=B21*RADIANS(B5-B13)', '0.000000'),
        (26, 'x of the station (m)', '=B24*SIN(B25)', '0.0'),
        (27, 'y of the station (m)', '=B23-B24*COS(B25)', '0.0'),
        (28, 'Grid spacing x (m)', '=B16-B15', '0.0'),
        (29, 'Grid spacing y (m)', '=B18-B17', '0.0'),
        (30, 'Grid column (fractional)', '=(B26-B15)/B28', '0.0000'),
        (31, 'Grid row (fractional)', '=(B27-B17)/B29', '0.0000'),
        (32, 'Column i of the grid points west of the station', '=INT(B30)', '0'),
        (33, 'Row j of the grid points south of the station', '=INT(B31)', '0'),
        (34, 'u (position between the columns, 0-1)', '=B30-B32', '0.0000'),
        (35, 'v (position between the rows, 0-1)', '=B31-B33', '0.0000'),
    ]:
        put(g, f'A{r}', lab)
        put(g, f'B{r}', f, fmt=fmt)
    put(g, 'A36', 'Column i, row j in HeatAQ\'s requests')
    put(g, 'B36', app['cell']['i'], f_app, '0')
    put(g, 'C36', app['cell']['j'], f_app, '0')
    put(g, 'D36', '=IF(AND(B32=B36,B33=C36),"OK","CHECK")', f_bold)

    put(g, 'A38', 'The 4 grid points', f_h2)
    link(g, 'C38', 'Open the grid points\' position, height and land ↗', raw['links']['sample_corners'])
    header(g, 39, ['Grid point', 'Row j', 'Column i', 'dj', 'di', 'Latitude', 'Longitude', 'Height (m)',
                   'Land fraction', 'Land?', 'Bilinear weight', 'Distance (grid cells)', 'Weight used',
                   'HeatAQ weight', 'Difference', 'Distance (km)'])
    g.row_dimensions[39].height = 30
    cn = raw['corners']
    for k, (label, (dj, di)) in enumerate(zip(['1 south-west', '2 south-east', '3 north-west', '4 north-east'],
                                              [(0, 0), (0, 1), (1, 0), (1, 1)])):
        r = 40 + k
        put(g, f'A{r}', label)
        put(g, f'B{r}', f'=$B$33+D{r}', fmt='0')
        put(g, f'C{r}', f'=$B$32+E{r}', fmt='0')
        put(g, f'D{r}', dj, fmt='0')
        put(g, f'E{r}', di, fmt='0')
        put(g, f'F{r}', cn['latitude'][dj][di], f_in, '0.00000')
        put(g, f'G{r}', cn['longitude'][dj][di], f_in, '0.00000')
        put(g, f'H{r}', cn['altitude'][dj][di], f_in, '0.0')
        put(g, f'I{r}', cn['land_area_fraction'][dj][di], f_in, '0.00')
        put(g, f'J{r}', f'=I{r}>=$B$46')
        put(g, f'K{r}', f'=IF(E{r}=1,$B$34,1-$B$34)*IF(D{r}=1,$B$35,1-$B$35)', fmt='0.0000')
        put(g, f'L{r}', f'=SQRT((E{r}-$B$34)^2+(D{r}-$B$35)^2)', fmt='0.0000')
        put(g, f'M{r}', f'=IF(OR($B$45=4,$B$45=0),K{r},IF(J{r},1/MAX(L{r},0.000001),0)/$B$47)', fmt='0.0000')
        put(g, f'N{r}', app['cell']['weights'][k], f_app, '0.0000')
        put(g, f'O{r}', f'=M{r}-N{r}', fmt='0.0000')
        put(g, f'P{r}', f'=6371*2*ASIN(SQRT(SIN(RADIANS(F{r}-$B$4)/2)^2+COS(RADIANS($B$4))*COS(RADIANS(F{r}))'
                        f'*SIN(RADIANS(G{r}-$B$5)/2)^2))', fmt='0.000')
    put(g, 'L44', 'Sum', f_bold)
    put(g, 'M44', '=SUM(M40:M43)', f_bold, '0.0000')
    for r, lab, val, fmt, note in [
        (45, 'Grid points on land', '=COUNTIF(J40:J43,TRUE)', '0', ''),
        (46, 'Land if land fraction ≥', LAND_THRESHOLD, '0.00', 'MetNordicForecast::LAND_THRESHOLD'),
        (47, 'Sum of 1/distance over land points', '=SUMPRODUCT(--J40:J43,1/L40:L43)', '0.0000', ''),
        (48, 'Method', '=IF(OR(B45=4,B45=0),"bilinear","inverse distance on land points")', None, ''),
        (49, 'Grid height, weighted (m)', '=SUMPRODUCT(M40:M43,H40:H43)', '0.0', ''),
        (50, 'Grid height minus station height (m)', '=ROUND(B49,1)-B6', '0.0',
         'HeatAQ rounds the grid height to 0.1 m first'),
        (51, 'Air temperature correction (°C)', '=B7*B50', '0.000', 'Added to every forecast air temperature'),
    ]:
        put(g, f'A{r}', lab)
        put(g, f'B{r}', val, f_in if r == 46 else f_norm, fmt, fill_in if r == 46 else None)
        put(g, f'D{r}', note, f_note)
    put(g, 'C49', app['cell']['grid_altitude'], f_app, '0.0')
    put(g, 'D49', 'HeatAQ (green)', f_note)
    put(g, 'A53', 'Check: the grid points\' coordinates, weighted, give back the station', f_h2)
    put(g, 'A54', 'Latitude from the grid points')
    put(g, 'B54', '=SUMPRODUCT(K40:K43,F40:F43)', fmt='0.00000')
    put(g, 'A55', 'Longitude from the grid points')
    put(g, 'B55', '=SUMPRODUCT(K40:K43,G40:G43)', fmt='0.00000')
    put(g, 'A56', 'Distance from the station (km)')
    put(g, 'B56', '=6371*2*ASIN(SQRT(SIN(RADIANS(B54-B4)/2)^2+COS(RADIANS(B4))*COS(RADIANS(B54))'
                  '*SIN(RADIANS(B55-B5)/2)^2))', fmt='0.000')
    put(g, 'C56', '=IF(B56<0.5,"OK","CHECK")', f_bold)
    put(g, 'D56', 'HeatAQ refuses a station whose position is off by more than 0.5 km', f_note)
    g.freeze_panes = 'A3'
    for name, ref in [('W1', 'Grid!$M$40'), ('W2', 'Grid!$M$41'), ('W3', 'Grid!$M$42'), ('W4', 'Grid!$M$43'),
                      ('AirCorrection', 'Grid!$B$51'), ('UtcOffset', 'Grid!$B$8')]:
        wb.defined_names[name] = DefinedName(name, attr_text=ref)

    # ------------------------------------------------------------ Runs
    runs = raw['runs']
    rs = wb.create_sheet('Runs')
    for col, w in zip('ABCDEFGHIJ', [18, 20, 20, 20, 8, 18, 14, 10, 10, 14]):
        rs.column_dimensions[col].width = w
    put(rs, 'A1', 'Forecast runs and when they became available', f_title)
    put(rs, 'A2', 'Available = the file\'s creation time ("Created" in its metadata) if it is 0-6 hours after the '
        'run time, else the run time + 1 hour. All times UTC unless marked local.', f_note)
    header(rs, 4, ['Run (UTC)', 'Created (file metadata)', 'Available (UTC)', 'Available (local)', 'Hours',
                   'Last hour (UTC)', 'Raw values', 'Metadata', 'File page', 'Day in catalog'])
    rs.row_dimensions[4].height = 30
    R0, R1 = 5, 5 + len(runs) - 1
    for k, run in enumerate(runs):
        r = R0 + k
        put(rs, f'A{r}', ts(run['issued']), f_in, DT)
        put(rs, f'B{r}', ts(run['created']) if run['created'] else None, f_in, DTS)
        put(rs, f'C{r}', f'=IF(AND(B{r}<>"",B{r}>=A{r},B{r}-A{r}<=6/24),B{r},A{r}+1/24)', fmt=DTS)
        put(rs, f'D{r}', f'=C{r}+UtcOffset/24', fmt=DTS)
        put(rs, f'E{r}', run['hours'], f_in, '0')
        put(rs, f'F{r}', f'=A{r}+(E{r}-1)/24', fmt=DT)
        link(rs, f'G{r}', 'Open data ↗', run['data'])
        link(rs, f'H{r}', 'Open ↗', run['das'])
        link(rs, f'I{r}', 'Open ↗', run['html'])
        link(rs, f'J{r}', 'Open ↗', run['catalog'])
    stored_avail = {row['issued']: row['available'] for row in app['rows']}
    put(rs, f'A{R1 + 2}', 'HeatAQ\'s available times:', f_bold)
    for k, run in enumerate(runs):
        r = R1 + 3 + k
        put(rs, f'A{r}', ts(run['issued']), f_in, DT)
        if run['issued'] in stored_avail:
            put(rs, f'B{r}', ts(stored_avail[run['issued']]), f_app, DTS)
            put(rs, f'C{r}', f'=IF(ABS(B{r}-C{R0 + k})<1/86400,"OK","CHECK")', f_bold)
        else:
            put(rs, f'B{r}', 'not read by HeatAQ', f_note)
    rs.freeze_panes = 'A5'
    RUNS_A, RUNS_C = f'Runs!$A${R0}:$A${R1}', f'Runs!$C${R0}:$C${R1}'
    RUNS_E, RUNS_F = f'Runs!$E${R0}:$E${R1}', f'Runs!$F${R0}:$F${R1}'

    # ------------------------------------------------------------ Forecast
    fc = wb.create_sheet('Forecast')
    put(fc, 'A1', 'Forecast values: recomputed from MET\'s raw values, and HeatAQ\'s values', f_title)
    put(fc, 'A2', 'Raw values (blue) are the 4 grid points 1-4 of sheet Grid. Air: kelvin; wind: m/s; humidity: '
        'fraction. Click a run time (first hour of each run) to open MET\'s raw values for that run.', f_note)
    FIRST = 7
    LAST = FIRST + sum(r['hours'] for r in runs) - 1
    put(fc, 'A3', 'Rows OK:', f_bold)
    put(fc, 'B3', f'=COUNTIF(Z{FIRST}:Z{LAST},"OK")', f_bold, '0')
    put(fc, 'C3', f'=" of "&COUNT(B{FIRST}:B{LAST})', f_bold)
    put(fc, 'E3', 'Largest difference:', f_bold)
    for lab_ref, val_ref, lab, col in [('G3', 'H3', 'air °C', 'K'), ('I3', 'J3', 'wind m/s', 'R'), ('K3', 'L3', 'RH %', 'Y')]:
        put(fc, lab_ref, lab)
        put(fc, val_ref, f'=MAX(MAX({col}{FIRST}:{col}{LAST}),-MIN({col}{FIRST}:{col}{LAST}))', fmt='0.000')
    for text, c1, c2 in [('Run and hour', 1, 4), ('Air temperature, grid points 1-4 (K)', 5, 8),
                         ('Air temperature (°C)', 9, 11), ('Wind speed, grid points 1-4 (m/s)', 12, 15),
                         ('Wind speed (m/s)', 16, 18), ('Relative humidity, grid points 1-4 (fraction)', 19, 22),
                         ('Relative humidity (%)', 23, 25), ('', 26, 27)]:
        fc.merge_cells(start_row=5, start_column=c1, end_row=5, end_column=c2)
        c = fc.cell(row=5, column=c1, value=text)
        c.font, c.fill, c.alignment = f_bold, fill_head, center
    header(fc, 6, ['Run (UTC)', 'Hour ahead', 'Valid (UTC)', 'Valid (local)']
           + ['1', '2', '3', '4', 'This sheet', 'HeatAQ', 'Difference'] * 3 + ['Check', 'Key'])
    for k, w in enumerate([17, 7, 17, 17] + ([9] * 4 + [10, 9, 10]) * 3 + [8, 7], start=1):
        fc.column_dimensions[get_column_letter(k)].width = w
    stored = {(row['issued'], row['valid']): row for row in app['rows']}
    r = FIRST
    for run in runs:
        iss = ts(run['issued'])
        for k in range(run['hours']):
            def c4(v):
                return [v[k][0][0], v[k][0][1], v[k][1][0], v[k][1][1]]
            sv = stored.get((run['issued'], (iss + dt.timedelta(hours=k)).strftime('%Y-%m-%d %H:%M:%S')))
            if k == 0:
                link(fc, f'A{r}', iss, run['data'], DT)
            else:
                put(fc, f'A{r}', iss, f_in, DT)
            put(fc, f'B{r}', k, f_in, '0')
            put(fc, f'C{r}', f'=A{r}+B{r}/24', fmt=DT)
            put(fc, f'D{r}', f'=C{r}+UtcOffset/24', fmt=DT)
            for cols, values, calc, key, fmt_raw, fmt_app in [
                ('EFGH', c4(run['air']), f'=E{r}*W1+F{r}*W2+G{r}*W3+H{r}*W4-273.15+AirCorrection', 'air', '0.000', '0.00'),
                ('LMNO', c4(run['wind']), f'=MAX(0,L{r}*W1+M{r}*W2+N{r}*W3+O{r}*W4)', 'wind', '0.000', '0.00'),
                ('STUV', c4(run['rh']), f'=MIN(100,MAX(0,100*(S{r}*W1+T{r}*W2+U{r}*W3+V{r}*W4)))', 'rh', '0.0000', '0.0'),
            ]:
                for col, v in zip(cols, values):
                    put(fc, f'{col}{r}', v, f_in, fmt_raw)
                c_calc = get_column_letter(fc[f'{cols[-1]}{r}'].column + 1)
                c_app = get_column_letter(fc[f'{cols[-1]}{r}'].column + 2)
                c_diff = get_column_letter(fc[f'{cols[-1]}{r}'].column + 3)
                put(fc, f'{c_calc}{r}', calc, fmt='0.000' if key != 'rh' else '0.00')
                put(fc, f'{c_app}{r}', float(sv[key]) if sv else None, f_app, fmt_app)
                put(fc, f'{c_diff}{r}', f'=IF({c_app}{r}="","",{c_calc}{r}-{c_app}{r})', fmt='0.000')
            put(fc, f'Z{r}', f'=IF(J{r}="","no HeatAQ value",IF(AND(ABS(K{r})<=0.0051,ABS(R{r})<=0.0051,'
                             f'ABS(Y{r})<=0.051),"OK","CHECK"))', f_bold)
            put(fc, f'AA{r}', f'=MATCH(A{r},{RUNS_A},0)*100+B{r}', f_note, '0')
            if k == run['hours'] - 1:
                for cidx in range(1, 28):
                    fc.cell(row=r, column=cidx).border = box
            r += 1
    fc.freeze_panes = f'E{FIRST}'
    put(fc, f'A{LAST + 2}', 'Tolerance: HeatAQ stores air and wind with 2 decimals (±0.005) and humidity with 1 '
        '(±0.05). Key: helper for sheet Night plan (run number × 100 + hour ahead).', f_note)
    FC_I, FC_KEY = f'Forecast!$I${FIRST}:$I${LAST}', f'Forecast!$AA${FIRST}:$AA${LAST}'

    # ------------------------------------------------------------ Night plan
    npl = wb.create_sheet('Night plan')
    for col, w in zip('ABCDEFGHIJ', [17, 17, 19, 10, 24, 14, 14, 14, 14, 14]):
        npl.column_dimensions[col].width = w
    put(npl, 'A1', 'Which forecast the predictive controller plans the night with', f_title)
    put(npl, 'A2', 'When the pool closes, the controller plans the night with the newest run available at that '
        'hour (and covering it). During the closed period it plans again whenever a newer run becomes available. '
        'Example opening hours below; change them (yellow) and the table follows.', f_note)
    npl.merge_cells('A2:I2')
    npl['A2'].alignment = wrap
    npl.row_dimensions[2].height = 40
    close_local = dt.datetime.combine(d0, dt.time.fromisoformat(args.close))
    open_local = dt.datetime.combine(d1, dt.time.fromisoformat(args.open))
    put(npl, 'A4', 'Pool closes (local)')
    put(npl, 'B4', close_local, f_in, DT, fill_in)
    put(npl, 'A5', 'Pool opens (local)')
    put(npl, 'B5', open_local, f_in, DT, fill_in)
    put(npl, 'A7', 'Each closed hour: the newest run, and what the controller does', f_h2)
    header(npl, 8, ['Hour (local)', 'Hour (UTC)', 'Newest run available (UTC)', 'Its age (h)', 'Controller',
                    'Air °C this hour (that run)'])
    npl.row_dimensions[8].height = 30
    T0, T1 = 9, 9 + 23

    def lookup(run_ref, hour_ref):
        """Air °C of run run_ref for hour hour_ref (UTC); hours past the run's end use its last hour."""
        n = f'INDEX({RUNS_E},MATCH({run_ref},{RUNS_A},0))'
        return (f'INDEX({FC_I},MATCH(MATCH({run_ref},{RUNS_A},0)*100+MIN({n}-1,ROUND(({hour_ref}-{run_ref})*24,0)),'
                f'{FC_KEY},0))')

    for k in range(24):
        r = T0 + k
        put(npl, f'A{r}', f'=IF($B$4+{k}/24<$B$5,$B$4+{k}/24,"")', fmt=DT)
        put(npl, f'B{r}', f'=IF(A{r}="","",A{r}-UtcOffset/24)', fmt=DT)
        put(npl, f'C{r}', f'=IF(A{r}="","",IFERROR(1/(1/_xlfn.MAXIFS({RUNS_A},{RUNS_C},"<="&B{r},{RUNS_F},">="&B{r})),'
                          f'"no run: actual weather"))', fmt=DT)
        put(npl, f'D{r}', f'=IF(OR(A{r}="",NOT(ISNUMBER(C{r}))),"",ROUND((B{r}-C{r})*24,1))', fmt='0.0')
        if k == 0:
            put(npl, f'E{r}', f'=IF(A{r}="","","Plans the night")', f_bold)
        else:
            put(npl, f'E{r}', f'=IF(A{r}="","",IF(C{r}<>C{r - 1},"Plans again (newer run)","Keeps the plan"))')
        put(npl, f'F{r}', f'=IF(OR(A{r}="",NOT(ISNUMBER(C{r}))),"",{lookup(f"C{r}", f"B{r}")})', fmt='0.00')

    # How the forecast for the same hour changes from run to run
    S0 = T1 + 3
    put(npl, f'A{S0}', 'Air temperature (°C) forecast for each hour, by run: how the forecast for the same hour '
        'changes', f_h2)
    close_utc = close_local - dt.timedelta(hours=utc_offset)
    open_utc = open_local - dt.timedelta(hours=utc_offset)
    shown = [run for run in runs if close_utc - dt.timedelta(hours=15) <= ts(run['issued']) <= open_utc + dt.timedelta(hours=6)]
    cols = [get_column_letter(3 + k) for k in range(len(shown))]
    for rr, lab in [(S0 + 1, 'Run (UTC)'), (S0 + 2, 'Used by the controller?')]:
        put(npl, f'A{rr}', lab, f_bold, fill=fill_head)
        put(npl, f'B{rr}', '', fill=fill_head)
    header(npl, S0 + 3, ['Hour (local)', 'Hour (UTC)'])
    for col, run in zip(cols, shown):
        put(npl, f'{col}{S0 + 1}', ts(run['issued']), f_in, DT, fill_head)
        put(npl, f'{col}{S0 + 2}', f'=IF(COUNTIF($C${T0}:$C${T1},{col}{S0 + 1})>0,"yes","no")', f_bold, fill=fill_head)
        put(npl, f'{col}{S0 + 3}', f'=TEXT({col}{S0 + 1},"dd hh")&"Z"', f_bold, fill=fill_head)
    H0 = S0 + 4
    n_hours = int((open_local - close_local).total_seconds() // 3600) + 7     # closed hours + 6 open hours
    for k in range(n_hours):
        r = H0 + k
        put(npl, f'A{r}', f'=$B$4+{k}/24', fmt=DT)
        put(npl, f'B{r}', f'=A{r}-UtcOffset/24', fmt=DT)
        for col in cols:
            rc = f'{col}${S0 + 1}'
            put(npl, f'{col}{r}', f'=IF(ROUND(($B{r}-{rc})*24,0)<0,"",{lookup(rc, f"$B{r}")})', fmt='0.00')
    put(npl, f'A{H0 + n_hours + 1}', 'Empty: the run was not issued yet for that hour. Hours past a run\'s last '
        'hour use its last hour (as the controller does). A plan looks at the closed hours and the open hours '
        'that follow.', f_note)
    npl.freeze_panes = 'A9'

    wb.save(out)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('--name', required=True, help='Station name (for the titles)')
    ap.add_argument('--lat', type=float, required=True)
    ap.add_argument('--lon', type=float, required=True)
    ap.add_argument('--elevation', type=float, required=True, help='Station height (m), as in HeatAQ')
    ap.add_argument('--date', required=True, help='First of the two days (UTC), YYYY-MM-DD, from 2019-11-28')
    ap.add_argument('--tz', default='Europe/Oslo', help='Local time zone (default Europe/Oslo)')
    ap.add_argument('--close', default='21:00', help='Pool closes on the first day, local time (default 21:00)')
    ap.add_argument('--open', default='07:00', help='Pool opens on the second day, local time (default 07:00)')
    ap.add_argument('--php', default='php', help='PHP command (default php)')
    ap.add_argument('--out', help='Output file (default: forecast_check_<name>_<date>.xlsx)')
    args = ap.parse_args()

    d0 = dt.date.fromisoformat(args.date)
    days = [d0, d0 + dt.timedelta(days=1)]
    close_local = dt.datetime.combine(d0, dt.time.fromisoformat(args.close))
    utc_offset = int(close_local.replace(tzinfo=ZoneInfo(args.tz)).utcoffset().total_seconds() // 3600)
    out = args.out or f"forecast_check_{re.sub(r'[^A-Za-z0-9]+', '_', args.name).strip('_')}_{args.date}.xlsx"

    log('HeatAQ\'s values (lib/MetNordicForecast.php)...')
    app = app_values(args.php, args.lat, args.lon, args.elevation, args.date)
    log(f'  {len(app["rows"])} forecast hours')
    log('Raw values from MET\'s archive...')
    raw = fetch_raw(args.lat, args.lon, days)
    build(args, raw, app, utc_offset, out)
    log(f'Saved {out}')


if __name__ == '__main__':
    main()
